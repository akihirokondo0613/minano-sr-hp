"""生成コピーの実式を評価して帳簿の境界条件を追試する。Office実計算ではない。

python3 scripts/shoshiki/test_excel_templates.py
原配布物は更新せず、TemporaryDirectory 内にだけ生成する。
"""
import datetime as dt
import calendar
import decimal
import functools
import json
import pathlib
import re
import tempfile
import zipfile

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries

import forms_base as B


class FormulaError(Exception):
    pass


@functools.lru_cache(maxsize=None)
def parse_formula(formula):
    """今回の配布式が使う構文だけを読む。Python evalやOfficeは使わない。"""
    tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
    index = 0
    priority = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1,
                "&": 2, "+": 3, "-": 3, "*": 4, "/": 4}

    def expression(level=0):
        nonlocal index
        token = tokens[index]
        index += 1
        if token.type == "FUNC" and token.subtype == "OPEN":
            args = []
            if tokens[index].subtype != "CLOSE":
                while True:
                    args.append(expression())
                    if tokens[index].type != "SEP":
                        break
                    index += 1
            assert tokens[index].subtype == "CLOSE"
            index += 1
            left = ("function", token.value[:-1].upper(), args)
        elif token.type == "PAREN" and token.subtype == "OPEN":
            left = expression()
            assert tokens[index].type == "PAREN" and tokens[index].subtype == "CLOSE"
            index += 1
        elif token.type == "OPERATOR-PREFIX":
            left = ("unary", token.value, expression(5))
        elif token.subtype == "NUMBER":
            left = ("value", float(token.value))
        elif token.subtype == "TEXT":
            left = ("value", token.value[1:-1].replace('""', '"'))
        elif token.subtype == "RANGE":
            left = ("reference", token.value.replace("$", ""))
        else:
            raise AssertionError(f"未対応トークン: {token}")
        while index < len(tokens):
            token = tokens[index]
            if token.type != "OPERATOR-INFIX" or priority.get(token.value, -1) < level:
                break
            index += 1
            left = ("operator", token.value, left, expression(priority[token.value] + 1))
        return left

    node = expression()
    assert index == len(tokens), formula
    return node


class FormulaReader:
    """保存された数式セルを限定的に評価する。Excel互換性全般の保証ではない。"""
    def __init__(self, ws):
        self.ws, self.cache = ws, {}

    def cell(self, address):
        if address not in self.cache:
            value = self.ws[address].value
            if isinstance(value, str) and value.startswith("="):
                value = self.evaluate(parse_formula(value))
            elif isinstance(value, dt.datetime):
                value = (value - dt.datetime(1899, 12, 30)).total_seconds() / 86400
            elif isinstance(value, dt.date):
                value = (value - dt.date(1899, 12, 30)).days
            elif isinstance(value, dt.time):
                value = (value.hour * 3600 + value.minute * 60 + value.second) / 86400
            self.cache[address] = value
        return self.cache[address]

    @staticmethod
    def number(value):
        return 0 if value in (None, "") else float(value)

    @staticmethod
    def same(a, b):
        if isinstance(a, str) and isinstance(b, str):
            return a.casefold() == b.casefold()
        return ("" if a is None else a) == ("" if b is None else b)

    @staticmethod
    def criterion_matches(value, criterion):
        if not isinstance(criterion, str):
            return FormulaReader.same(value, criterion)
        if criterion.startswith("="): criterion = criterion[1:]
        pattern, index = "", 0
        while index < len(criterion):
            char = criterion[index]
            if char == "~" and index + 1 < len(criterion):
                index += 1
                pattern += re.escape(criterion[index])
            else:
                pattern += ".*" if char == "*" else "." if char == "?" else re.escape(char)
            index += 1
        return re.fullmatch(pattern, "" if value is None else str(value), re.IGNORECASE) is not None

    def evaluate(self, node):
        kind = node[0]
        if kind == "value":
            return node[1]
        if kind == "reference":
            address = node[1]
            if ":" not in address:
                return self.cell(address)
            c1, r1, c2, r2 = range_boundaries(address)
            return [self.cell(self.ws.cell(r, c).coordinate)
                    for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        if kind == "unary":
            value = self.number(self.evaluate(node[2]))
            return -value if node[1] == "-" else value
        if kind == "operator":
            op = node[1]
            a, b = self.evaluate(node[2]), self.evaluate(node[3])
            if op == "&":
                return str(a) + str(b)
            if op in ("=", "<>"):
                equal = self.same(a, b)
                return equal if op == "=" else not equal
            a, b = self.number(a), self.number(b)
            return {"+": lambda: a + b, "-": lambda: a - b, "*": lambda: a * b,
                    "/": lambda: a / b, "<": lambda: a < b, ">": lambda: a > b,
                    "<=": lambda: a <= b, ">=": lambda: a >= b}[op]()
        name, arguments = node[1:]
        if name == "IF":
            return self.evaluate(arguments[1 if self.evaluate(arguments[0]) else 2])
        if name == "NA":
            raise FormulaError("#N/A")
        values = [self.evaluate(x) for x in arguments]
        flat = [x for value in values for x in (value if isinstance(value, list) else [value])]
        numbers = [x for x in flat if isinstance(x, (float, int))]
        if name == "OR": return any(values)
        if name == "AND": return all(values)
        if name == "COUNT": return len(numbers)
        if name == "COUNTA": return sum(x is not None for x in flat)
        if name == "SUM": return sum(numbers)
        if name == "MAX": return max(numbers, default=0)
        if name == "MIN": return min(numbers, default=0)
        if name == "TIME": return (values[0] * 3600 + values[1] * 60 + values[2]) % 86400 / 86400
        if name == "ROUND":
            return float(decimal.Decimal(str(values[0])).quantize(
                decimal.Decimal(10) ** -int(values[1]), rounding=decimal.ROUND_HALF_UP))
        if name == "SUMIFS":
            return sum(self.number(value) for i, value in enumerate(values[0])
                       if all(self.criterion_matches(values[j][i], values[j + 1])
                              for j in range(1, len(values), 2)))
        if name == "SUBSTITUTE": return values[0].replace(values[1], values[2])
        if name == "DATEVALUE":
            date = dt.date(*map(int, values[0].split("/")))
            return (date - dt.date(1899, 12, 30)).days
        if name in ("DAY", "YEAR", "MONTH"):
            date = dt.date(1899, 12, 30) + dt.timedelta(days=values[0])
            return {"DAY": date.day, "YEAR": date.year, "MONTH": date.month}[name]
        if name == "EOMONTH":
            date = dt.date(1899, 12, 30) + dt.timedelta(days=values[0])
            month = date.year * 12 + date.month - 1 + int(values[1])
            year, zero_month = divmod(month, 12)
            date = dt.date(year, zero_month + 1, calendar.monthrange(year, zero_month + 1)[1])
            return (date - dt.date(1899, 12, 30)).days
        if name == "DATE":
            year, month, day = map(int, values)
            date = dt.date(year, month, 1) + dt.timedelta(days=day - 1)
            return (date - dt.date(1899, 12, 30)).days
        if name == "TEXT" and values[1] == "aaa":
            date = dt.date(1899, 12, 30) + dt.timedelta(days=values[0])
            return "月火水木金土日"[date.weekday()]
        raise AssertionError(f"未対応関数: {name}")


def run():
    checks = []
    with tempfile.TemporaryDirectory(prefix="shoshiki-excel-test-") as temp:
        paths = {no: pathlib.Path(temp) / f"{no}.xlsx" for no in ("D-31", "D-32")}
        B.build_shukkinbo(paths["D-31"])
        B.build_yukyu(paths["D-32"])

        def check(no, name, inputs, expected):
            ws = openpyxl.load_workbook(paths[no]).active
            for address, value in inputs.items(): ws[address] = value
            reader = FormulaReader(ws)
            actual = {}
            for address, wanted in expected.items():
                try: result = reader.cell(address)
                except FormulaError as error: result = str(error)
                assert result == wanted, (name, address, wanted, result)
                actual[address] = result
            checks.append({"template": no, "case": name, "actual": actual})

        shifts = [
            ("通常勤務", "09:00", "18:00", 60, 0, 8, 0),
            ("翌日勤務", "21:00", "03:00", 0, 0, 6, 5),
            ("早朝勤務", "00:00", "04:00", 0, 0, 4, 4),
            ("深夜休憩を別途控除", "23:00", "07:00", 60, 30, 7, 5.5),
            ("深夜中だけの勤務と休憩", "00:00", "04:00", 60, 60, 3, 3),
            ("23時間半勤務の翌22時", "23:00", "22:30", 0, 0, 23.5, 6.5),
            ("同時刻は0時間", "09:00", "09:00", 0, 0, 0, 0),
        ]
        for name, start, end, rest, deep_rest, work, deep in shifts:
            time = lambda x: dt.time(*map(int, x.split(":")))
            check("D-31", name, {"D6": time(start), "E6": time(end), "F6": rest, "K6": deep_rest},
                  {"G6": work, "I6": deep})
        check("D-31", "27時入力も6時間", {"D6": dt.time(21), "E6": 27 / 24, "F6": 0, "K6": 0}, {"G6": 6, "I6": 5})
        check("D-31", "深夜休憩が総休憩を超える", {"D6": dt.time(21), "E6": dt.time(3), "F6": 0, "K6": 30}, {"I6": "#N/A", "I37": "#N/A"})
        check("D-31", "日勤で深夜休憩を入力", {"D6": dt.time(9), "E6": dt.time(18), "F6": 60, "K6": 30}, {"I6": "#N/A"})
        check("D-31", "総休憩が勤務を超える", {"D6": dt.time(9), "E6": dt.time(10), "F6": 120}, {"G6": "#N/A", "G37": "#N/A"})
        check("D-31", "24時間勤務は分割記録が必要", {"D6": 0, "E6": 1, "F6": 0}, {"G6": "#N/A", "I6": "#N/A"})
        check("D-31", "負の始業時刻はエラー", {"D6": -0.25, "E6": 0, "F6": 0}, {"G6": "#N/A", "I6": "#N/A"})
        check("D-31", "未入力は空欄", {}, {"G6": "", "H6": "", "I6": ""})
        for month, days in (("2026/02", 28), ("2028/02", 29), ("2026/11", 30), ("2026/05", 31)):
            expected = {f"A{days + 5}": days}
            if days < 31: expected.update({f"{column}{days + 6}": "" for column in "ABCFGHI"})
            check("D-31", f"月末境界{month}", {"H2": month}, expected)
        check("D-31", "月末超過行の古い勤務を集計しない", {"H2": "2026/02", "D34": dt.time(9), "E34": dt.time(18)}, {"A34": "", "G34": "", "G37": 0})
        check("D-31", "日勤と夜勤の月計", {"D6": dt.time(9), "E6": dt.time(18), "D7": dt.time(21), "E7": dt.time(3), "F7": 0}, {"G37": 14, "I37": 5})

        base = {"A6": "EMP001", "B6": "監査用架空社員", "D6": dt.date(2026, 4, 1), "E6": 0, "F6": 20,
                "A7": "EMP001", "B7": "監査用架空社員", "D7": dt.date(2026, 4, 1), "L7": dt.date(2026, 5, 16)}
        base.update({f"{openpyxl.utils.get_column_letter(12 + i)}6": dt.date(2026, 5, i + 1) for i in range(15)})
        check("D-32", "16日取得の継続行合算", base, {"G6": 20, "H6": 16, "I6": 4, "J6": "済", "G7": 20, "H7": 16, "I7": 4, "J7": "済"})
        check("D-32", "同姓の別IDを分ける", {**base, "A7": "EMP002", "E7": 0, "F7": 10}, {"H6": 15, "I6": 5, "J6": "済", "H7": 1, "I7": 9, "J7": "未達"})
        check("D-32", "IDの先頭ゼロを区別", {**base, "A6": "001", "A7": "1", "F7": 10}, {"H6": 15, "I6": 5, "H7": 1, "I7": 9})
        check("D-32", "IDのワイルドカードを文字として照合", {**base, "A6": "EMP*", "A7": "EMP001", "F7": 10}, {"H6": 15, "I6": 5, "H7": 1, "I7": 9})
        check("D-32", "疑問符とチルダを含む同一IDを合算", {**base, "A6": "EMP~?", "A7": "EMP~?"}, {"H6": 16, "I6": 4, "H7": 16, "I7": 4})
        check("D-32", "同IDでも異なる基準日は分ける", {**base, "D7": dt.date(2027, 4, 1), "E7": 0, "F7": 20}, {"H6": 15, "I6": 5, "H7": 1, "I7": 19, "J7": "未達"})
        check("D-32", "付与の二重入力は警告", {**base, "E7": 0, "F7": 20}, {"G6": "", "I6": "", "J6": "付与は1行に入力", "G7": "", "J7": "付与は1行に入力"})
        check("D-32", "社員ID未入力は合算しない", {"B6": "監査用架空社員", "D6": dt.date(2026, 4, 1), "F6": 20}, {"G6": "", "H6": "", "J6": "ID・基準日要入力"})
        check("D-32", "基準日未入力は合算しない", {"A6": "EMP001", "F6": 20}, {"G6": "", "H6": "", "J6": "ID・基準日要入力"})
        check("D-32", "付与入力なしを警告", {"A6": "EMP001", "D6": dt.date(2026, 4, 1)}, {"G6": "", "J6": "付与は1行に入力"})
        check("D-32", "半日手修正も継続行へ反映", {**base, "AA6": 2.5, "AA7": 1.5}, {"H6": 4, "I6": 16, "J6": "未達", "H7": 4, "I7": 16, "J7": "未達"})
        check("D-32", "半日を含む5日達成", {**base, "AA6": 2.5, "AA7": 2.5}, {"H6": 5, "I6": 15, "J6": "済", "J7": "済"})
        check("D-32", "繰越があっても今年付与10日未満は対象外", {"A6": "EMP001", "D6": dt.date(2026, 4, 1), "E6": 10, "F6": 5}, {"G6": 15, "J6": "対象外"})
        check("D-32", "空のテンプレート", {}, {"G6": "", "H6": "", "I6": "", "J6": ""})

        for path in paths.values():
            with zipfile.ZipFile(path) as package: assert package.testzip() is None
            ws = openpyxl.load_workbook(path).active
            assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_A4)
            assert ws.print_title_rows == "$1:$5"
        ws = openpyxl.load_workbook(paths["D-32"]).active
        assert [ws.cell(5, col).value for col in range(12, 27)] == [f"取得日{i}" for i in range(1, 16)]
        assert ws.column_dimensions["AB"].hidden
        assert ws.column_dimensions["AC"].hidden
    print(json.dumps({"passed": True, "cases": len(checks), "scope": "生成コピーの保存済み数式を限定評価。Office実計算・印刷は未検証。", "checks": checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    run()
